import os
import extract_msg
import re
import openpyxl


WORKLOAD_MANAGEMENT_PATH = r"C:\Users\RSPARLIN\HC-SC PHAC-ASPC\Bureau of Licensing Services and Systems (NNHPD) - Transitory Administration\Workload Management GC Forms"
SUBMISSION_PATTERN = r"Submission Number\s*(\d+)"

if __name__ == '__main__':
    wb = openpyxl.workbook.Workbook()
    ws = wb.active
    row = 2

    ws.cell(row=1, column=1).value = "SUBMISSION"
    ws.cell(row=1, column=2).value = "SOURCE"

    count = 0
    for root, dirs, files in os.walk(WORKLOAD_MANAGEMENT_PATH):
        for filename in files:
            if str(filename).lower().endswith('.msg'):
                count += 1
                # Construct the full path of the file
                msg_path = os.path.join(root, filename)

                try:
                    with extract_msg.Message(msg_path) as msg:
                        # Search for the submission number in the email body
                        submission_numbers = re.findall(SUBMISSION_PATTERN, msg.body, re.IGNORECASE)
                        print(f"Submission Numbers: {', '.join(submission_numbers)}")

                        for sub in submission_numbers:
                            ws.cell(row=row, column=1).value = sub if len(sub) == 6 else ""
                            ws.cell(row=row, column=2).value = msg_path
                            row += 1

                except Exception as e:
                    print(f"Could not process {msg_path}. Error: {e}")

                print(count)
                print("*" * 50)

    wb.save("map.xlsx")
    wb.close()