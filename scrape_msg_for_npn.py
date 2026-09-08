import os
import extract_msg
import re
import openpyxl



PATH = r"C:\Development\CMS\Discontinuations_AUG 2026"


if __name__ == '__main__':

    wb = openpyxl.workbook.Workbook()
    ws = wb.active
    row = 2

    ws.cell(row=1, column=1).value = "NPN"
    ws.cell(row=1, column=2).value = "SOURCE"

    count = 0

    npn_dict = {}
    files_dict = {}

    for root, dirs, files in os.walk(PATH):

        for filename in files:

            company_number = filename.split("-", 1)[0]
            npn_dict.setdefault(company_number, [])
            files_dict.setdefault(company_number, [])

            if str(filename).lower().endswith('.msg'):
                msg_path = os.path.join(root, filename)
                files_dict[company_number].append(msg_path)

                try:
                    with extract_msg.Message(msg_path) as msg:
                        # Search for the NPN in the email body
                        matches = re.findall(r'\b\d{8}\b', msg.body)

                        for match in matches:
                            if match not in npn_dict[company_number]:
                                npn_dict[company_number].append(match)

                except Exception as e:
                    print(f"Could not process {msg_path}. Error: {e}")

    for k, v in npn_dict.items():

        files_len = len(files_dict[k])

        for npn in v:
            for i in range(files_len):

                ws.cell(row=row, column=1).value = npn
                ws.cell(row=row, column=2).value = files_dict[k][i]
                row += 1

    wb.save("test.xlsx")


